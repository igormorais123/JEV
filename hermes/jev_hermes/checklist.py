"""Checklist: um documento, uma lista fixa de perguntas, um semáforo por item — numa chamada só.

    cd /root/.hermes/integrations/jev
    python3 -m jev_hermes.checklist --documento contrato.pdf --lista contrato-prestacao-de-servicos
    python3 -m jev_hermes.checklist --documento - --lista caminho/da/minha-lista.json --json

Porte para o Hermes da camada `integracao/camadas/checklist.py` do repositório JEV (R50), com
leitura de .pdf (pdftotext) e .docx além de texto.

Para que serve: triagem de documento contra pontos de atenção que se repetem (contrato, edital,
proposta, política). O Jev não escreve nem resume; ele responde cada pergunta da lista e informa uma
probabilidade estimada, que não garante acerto. O que ele deixa em amarelo é o que uma pessoa, ou o modelo caro, precisa
ler — A economia de leitura precisa ser medida no fluxo completo.

Por que uma chamada só: todas as perguntas leem o mesmo documento, e o contrato aceita várias
perguntas no mesmo payload. Medido: com 4, 16 e 64 perguntas a latência é a mesma de uma (R33,
R49), e a resposta principal se mantém em 96–98%. Mandar item por item multiplica a espera pelo
número de itens e não melhora o acerto (R50).

A lista é um JSON em `jev_hermes/listas/` (ou um caminho): `itens` com `id`, `tipo`
(`choice` ou `noul`), `pergunta`, `opcoes` (rótulo → descrição, só em `choice`) e `risco` (os
rótulos que acendem o vermelho; em `noul`, `["sim"]` ou `["nao"]`). Regras que o estudo mediu e
que esta camada aplica sozinha:

- Toda pergunta `choice` ganha a opção `nao-consta`. Sem classe de escape o Jev inventa uma
  resposta e vem confiante (R13).
- O corte é na **probabilidade** da resposta (0,90), não na confiança: a confiança depende do
  número de opções (R32) e itens da mesma lista têm números diferentes de opções.
- Pergunte o que o documento **estabelece**, não de que assunto ele fala: o Jev acha o tema
  quase sempre e erra quando o tema aparece sem a coisa (R28, R31).
- Um sentinela vai junto. Se o documento tenta dar ordem a quem o classifica, o resultado
  inteiro sai marcado e nada fica verde.
- Falha de rede, teto ou contrato devolve "não verificado"; nunca verde.
"""
import argparse
import hashlib
import json
import math
import subprocess
import sys
from pathlib import Path

from jev_hermes import camadas, nucleo

LISTAS = Path(__file__).resolve().parent / 'listas'
CORTE = 0.90
LIMITE_DO_DOCUMENTO = 60000  # caracteres; cerca de 20 mil tokens, abaixo do contexto publicado de 32 mil
NAO_CONSTA = 'O documento nao trata deste ponto.'


def carregar_lista(nome):
    caminho = Path(nome)
    if not caminho.exists():
        caminho = LISTAS / (nome if nome.endswith('.json') else nome + '.json')
    lista = json.loads(caminho.read_text(encoding='utf-8'))
    validar_lista(lista)
    return lista


def validar_lista(lista):
    """Rejeita rubricas ambíguas antes de qualquer transmissão."""
    if not isinstance(lista, dict) or not isinstance(lista.get('nome'), str) or not lista['nome'].strip():
        raise ValueError('lista precisa de nome')
    if not isinstance(lista.get('itens'), list) or not lista['itens']:
        raise ValueError('lista precisa de itens')
    ids = set()
    ativos = 0
    for item in lista['itens']:
        if not isinstance(item, dict):
            raise ValueError('item inválido')
        ident = item.get('id')
        if not isinstance(ident, str) or not ident.strip() or ident.startswith('_') or ident in ids:
            raise ValueError('id vazio, reservado ou duplicado')
        ids.add(ident)
        if not isinstance(item.get('ativo', True), bool):
            raise ValueError('ativo deve ser booleano')
        ativos += item.get('ativo', True)
        if not isinstance(item.get('pergunta'), str) or not item['pergunta'].strip():
            raise ValueError('pergunta vazia')
        if not isinstance(item.get('risco'), list) or not all(isinstance(x, str) for x in item['risco']):
            raise ValueError('risco deve ser lista de rótulos')
        if item.get('tipo') not in ('choice', 'noul'):
            raise ValueError(f"item {item['id']}: tipo deve ser choice ou noul")
        if item['tipo'] == 'choice':
            opcoes = item.get('opcoes')
            if not isinstance(opcoes, dict) or len(opcoes) < 2 or 'nao-consta' in opcoes or not all(isinstance(k, str) and k.strip() and isinstance(v, str) and v.strip() for k, v in opcoes.items()):
                raise ValueError('opções inválidas; nao-consta é reservado')
        if not set(item['risco']) <= (set(item['opcoes']) if item['tipo'] == 'choice' else {'sim', 'nao'}):
            raise ValueError(f"item {item['id']}: risco cita opção que não existe")
    if not ativos:
        raise ValueError('lista sem itens ativos')


def perguntas_da(lista):
    validar_lista(lista)
    perguntas = {}
    for item in lista['itens']:
        if not item.get('ativo', True):
            continue
        if item['tipo'] == 'noul':
            perguntas[item['id']] = {'type': 'noul', 'instructions': item['pergunta']}
        else:
            perguntas[item['id']] = {'type': 'choice', 'instructions': item['pergunta'] + ' Use apenas o texto do documento.',
                                     'criteria': {**item['opcoes'], 'nao-consta': NAO_CONSTA}}
    perguntas['_sentinela'] = dict(camadas.SENTINELA['sentinela'])
    return perguntas


def ler_item(item, bloco):
    """Devolve (resposta, probabilidade) de um item, qualquer que seja o tipo."""
    if not isinstance(bloco, dict) or not bloco:
        return None, None
    if item['tipo'] == 'noul':
        p = bloco.get('noul')
        return (None, None) if not probabilidade_valida(p) else (('sim', p) if p >= 0.5 else ('nao', 1 - p))
    escolha = bloco.get('choice')
    if escolha not in {*item['opcoes'], 'nao-consta'}:
        return None, None
    probabilidades = bloco.get('probabilities')
    p = probabilidades.get(escolha) if isinstance(probabilidades, dict) else None
    if p is None and bloco.get('confidence') is not None:  # a confiança é (K·p − 1)/(K − 1)
        if not probabilidade_valida(bloco['confidence']):
            return None, None
        k = len(item['opcoes']) + 1
        p = bloco['confidence'] * (k - 1) / k + 1 / k
    return (escolha, p) if probabilidade_valida(p) else (None, None)


def probabilidade_valida(valor):
    return type(valor) in (int, float) and math.isfinite(valor) and 0 <= valor <= 1


def semaforo(item, resposta, p, suspeito):
    if resposta is None:
        return 'não verificado'
    if suspeito or (p or 0) < CORTE:
        return 'amarelo'
    if resposta == 'nao-consta':
        return 'não consta'
    return 'vermelho' if resposta in item['risco'] else 'verde'


def auditar(documento, lista, *, transporte=None, tempo_total=30.0):
    validar_lista(lista)
    if not isinstance(documento, str) or not documento.strip() or len(documento) > LIMITE_DO_DOCUMENTO:
        raise ValueError('documento vazio ou acima do limite; não foi enviado')
    resultados = nucleo.classificar_em_paralelo(
        [f'DOCUMENTO:\n{documento}'], perguntas_da(lista), origem='camada-checklist',
        tempo_total=tempo_total, transporte=transporte, limite=LIMITE_DO_DOCUMENTO)
    respostas, detalhe = resultados[0]
    vigia = ((respostas or {}).get('_sentinela') or {}).get('choice')
    suspeito = vigia != 'nao-tenta'
    itens = []
    for item in lista['itens']:
        if not item.get('ativo', True):
            continue
        resposta, p = ler_item(item, (respostas or {}).get(item['id']))
        itens.append({'id': item['id'], 'pergunta': item['pergunta'], 'resposta': resposta,
                      'probabilidade': None if p is None else round(p, 3), 'cor': semaforo(item, resposta, p, suspeito)})
    cores = [i['cor'] for i in itens]
    return {'lista': lista['nome'], 'lista_snapshot': lista,
            'lista_sha256': hashlib.sha256(json.dumps(lista, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
            'documento_sha256': hashlib.sha256(documento.encode()).hexdigest(),
            'revisao_humana_obrigatoria': True,
            'caracteres': len(documento), **nucleo.resumo_das_chamadas(resultados),
            'latencia_ms': (detalhe or {}).get('latencia_ms'), 'tenta_instruir': suspeito, 'itens': itens,
            **{cor.replace(' ', '_').replace('ã', 'a'): cores.count(cor) for cor in ('vermelho', 'amarelo', 'verde', 'não consta', 'não verificado')}}


def ler_documento(caminho):
    """Texto de .pdf (pdftotext -layout), .docx (python-docx) ou arquivo de texto."""
    caminho = Path(caminho)
    sufixo = caminho.suffix.lower()
    if sufixo == '.pdf':
        return subprocess.run(['pdftotext', '-layout', str(caminho), '-'], capture_output=True, text=True,
                              timeout=120, check=True).stdout
    if sufixo == '.docx':
        import docx
        documento = docx.Document(str(caminho))
        partes = [p.text for p in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                partes.append(' | '.join(c.text for c in linha.cells))
        return '\n'.join(partes)
    return caminho.read_text(encoding='utf-8', errors='replace')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--documento', required=True, help='arquivo de texto, ou - para a entrada padrão')
    parser.add_argument('--lista', required=True, help='nome de uma lista em jev_hermes/listas/ ou caminho de um JSON')
    parser.add_argument('--json', action='store_true')
    args = parser.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except (AttributeError, ValueError):
        pass
    documento = sys.stdin.read() if args.documento == '-' else ler_documento(args.documento)
    if len(documento) > LIMITE_DO_DOCUMENTO:
        parser.error(f'documento com {len(documento)} caracteres; o máximo é {LIMITE_DO_DOCUMENTO}. Divida por capítulo.')
    resultado = auditar(documento, carregar_lista(args.lista))
    camadas.registrar('checklist', **{k: v for k, v in resultado.items() if k not in ('itens', 'lista_snapshot')},
                      respostas=[(i['id'], i['resposta'], i['probabilidade'], i['cor']) for i in resultado['itens']])
    if args.json:
        print(json.dumps(resultado, ensure_ascii=False))
        return 0
    print(f"[jev/checklist] {resultado['lista']} — {resultado['caracteres']} caracteres, {resultado['chamadas']} chamada(s), "
          f"US$ {nucleo.dec(resultado['custo_usd'], 6)}: {resultado['vermelho']} vermelho(s), {resultado['amarelo']} amarelo(s), "
          f"{resultado['verde']} verde(s), {resultado['nao_consta']} não consta.")
    if resultado['tenta_instruir']:
        print('ATENÇÃO: o documento contém texto que tenta dar ordens a quem o classifica. Nada foi marcado verde; leia você.')
    for i in resultado['itens']:
        p = nucleo.dec(i['probabilidade']) if i['probabilidade'] is not None else '—'
        print(f"- {i['cor'].upper()} ({p}) {i['id']}: {i['resposta']}")
    print('Amarelo quer dizer "leia você": a probabilidade ficou abaixo de 0,90. Verde não é parecer: é triagem.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
